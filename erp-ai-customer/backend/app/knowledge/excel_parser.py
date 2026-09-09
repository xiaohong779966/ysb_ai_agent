"""Parse and validate the first worksheet of an ERP FAQ XLSX workbook."""

from io import BytesIO
from pathlib import Path
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.exceptions import InvalidFileException

from backend.app.knowledge.exceptions import (
    InvalidKnowledgeHeadersError,
    InvalidWorkbookError,
    UnsupportedWorkbookTypeError,
)
from backend.app.knowledge.models import (
    KnowledgeParseResult,
    KnowledgeRecord,
    KnowledgeValidationIssue,
)

REQUIRED_HEADERS = ("问题", "答案")
XLSX_SUFFIX = ".xlsx"


def parse_knowledge_workbook(path: str | Path) -> KnowledgeParseResult:
    """Parse an XLSX workbook from disk without modifying it."""
    workbook_path = Path(path)
    _validate_filename(workbook_path.name)
    try:
        with workbook_path.open("rb") as workbook_file:
            return parse_knowledge_stream(workbook_file, filename=workbook_path.name)
    except OSError as exc:
        raise InvalidWorkbookError(f"无法读取知识库文件：{workbook_path}") from exc


def parse_knowledge_bytes(content: bytes, *, filename: str) -> KnowledgeParseResult:
    """Parse an uploaded XLSX payload held in memory."""
    return parse_knowledge_stream(BytesIO(content), filename=filename)


def parse_knowledge_stream(
    stream: BinaryIO,
    *,
    filename: str,
) -> KnowledgeParseResult:
    """Parse an XLSX binary stream and return valid records plus row issues."""
    _validate_filename(filename)
    try:
        workbook = load_workbook(
            stream,
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise InvalidWorkbookError("文件不是有效的 XLSX 工作簿或文件已损坏") from exc

    try:
        if not workbook.worksheets:
            raise InvalidWorkbookError("XLSX 工作簿不包含可读取的工作表")
        worksheet = workbook.worksheets[0]
        _validate_headers(worksheet)
        return _parse_worksheet(worksheet)
    finally:
        workbook.close()


def _validate_filename(filename: str) -> None:
    if Path(filename).suffix.lower() != XLSX_SUFFIX:
        raise UnsupportedWorkbookTypeError("仅支持 .xlsx 格式的知识库文件")


def _validate_headers(worksheet: object) -> None:
    first_header = _normalize_cell_value(worksheet.cell(row=1, column=1).value)
    second_header = _normalize_cell_value(worksheet.cell(row=1, column=2).value)
    extra_headers = [
        _normalize_cell_value(worksheet.cell(row=1, column=column).value)
        for column in range(3, worksheet.max_column + 1)
    ]

    if (first_header, second_header) != REQUIRED_HEADERS or any(extra_headers):
        actual_headers = tuple(
            _normalize_cell_value(worksheet.cell(row=1, column=column).value)
            for column in range(1, worksheet.max_column + 1)
        )
        raise InvalidKnowledgeHeadersError(
            "知识库表头必须严格为两列：问题、答案；"
            f"当前表头为：{actual_headers or ('<空>',)}"
        )


def _parse_worksheet(worksheet: object) -> KnowledgeParseResult:
    records: list[KnowledgeRecord] = []
    issues: list[KnowledgeValidationIssue] = []
    seen_questions: dict[str, int] = {}
    total_rows = max(worksheet.max_row - 1, 0)

    for row_number in range(2, worksheet.max_row + 1):
        question_cell: Cell = worksheet.cell(row=row_number, column=1)
        answer_cell: Cell = worksheet.cell(row=row_number, column=2)
        question = _normalize_cell_value(question_cell.value)
        answer = _normalize_cell_value(answer_cell.value)
        row_issues: list[KnowledgeValidationIssue] = []

        if question_cell.data_type == "f":
            row_issues.append(
                _issue(row_number, "question", "formula_not_allowed", "问题不能使用公式")
            )
        if answer_cell.data_type == "f":
            row_issues.append(
                _issue(row_number, "answer", "formula_not_allowed", "答案不能使用公式")
            )

        if not question and not answer:
            row_issues.append(_issue(row_number, "row", "empty_row", "问题和答案均为空"))
        else:
            if not question:
                row_issues.append(
                    _issue(row_number, "question", "required", "问题不能为空")
                )
            if not answer:
                row_issues.append(_issue(row_number, "answer", "required", "答案不能为空"))

        question_key = _question_key(question)
        if question:
            if question_key in seen_questions:
                first_row = seen_questions[question_key]
                row_issues.append(
                    _issue(
                        row_number,
                        "question",
                        "duplicate",
                        f"问题与第 {first_row} 行重复",
                    )
                )
            else:
                seen_questions[question_key] = row_number

        if row_issues:
            issues.extend(row_issues)
            continue

        records.append(
            KnowledgeRecord(
                row_number=row_number,
                question=question,
                answer=answer,
            )
        )

    return KnowledgeParseResult(
        sheet_name=worksheet.title,
        total_rows=total_rows,
        records=tuple(records),
        issues=tuple(issues),
    )


def _normalize_cell_value(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _question_key(question: str) -> str:
    return " ".join(question.split()).casefold()


def _issue(
    row_number: int,
    field: str,
    code: str,
    message: str,
) -> KnowledgeValidationIssue:
    return KnowledgeValidationIssue(
        row_number=row_number,
        field=field,
        code=code,
        message=message,
    )
